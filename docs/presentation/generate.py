"""Build the editable Russian talk deck; all measurements use repository data.

From the repository root:
  python -m pip install --target .cache/presentation/python python-pptx
  python docs/presentation/generate.py

Dependencies and QA renders stay in .cache/presentation, outside app manifests.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.cache/presentation/python'))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = ROOT / 'DunCarBox_Презентация.pptx'
W, H = 13.333333, 7.5
NAVY = '172D3C'
DEEP = '102331'
INK = '183342'
MUTED = '61727C'
TEAL = '087F8C'
MINT = 'A7E4DA'
CORAL = 'EF765C'
GOLD = 'EAB65A'
BG = 'F3F6F5'
WHITE = 'FFFFFF'
LINE = 'DCE5E4'
PALE_TEAL = 'E4F2EF'
PALE_CORAL = 'FCEBE5'
FONT = 'Segoe UI'

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = 'DunCarBox — от заказа к понятной укладке'
prs.core_properties.subject = 'Презентация хакатонного MVP, 5–7 минут'
prs.core_properties.author = 'DunCarBox'
prs.core_properties.keywords = 'упаковка, склад, 3D, Z3, FastAPI, React'
prs.core_properties.comments = 'Редактируемые тексты и схемы. Демо-метрики взяты из demo/*.response.json.'


def color(value):
    return RGBColor.from_string(value)


def shape(slide, x, y, w, h, fill, border=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = color(fill)
    if border:
        item.line.color.rgb = color(border)
        item.line.width = Pt(1)
    else:
        item.line.fill.background()
    if radius:
        item.adjustments[0] = 0.12
    return item


def text(slide, value, x, y, w, h, size=24, fill=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, spacing=1.12):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = item.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(fill)
    return item


def line(slide, x1, y1, x2, y2, fill=LINE, width=1.2):
    item = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    item.line.color.rgb = color(fill)
    item.line.width = Pt(width)
    return item


def circle(slide, x, y, d, fill):
    item = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    item.fill.solid()
    item.fill.fore_color.rgb = color(fill)
    item.line.fill.background()
    return item


def polygon(slide, pts, fill, border=None):
    # Freeform vertices round to integers before scaling: use 1/1000-inch units.
    vertices = [(round(x*1000), round(y*1000)) for x, y in pts]
    builder = slide.shapes.build_freeform(*vertices[0], scale=Inches(1)/1000)
    builder.add_line_segments(vertices[1:], close=True)
    item = builder.convert_to_shape()
    item.fill.solid()
    item.fill.fore_color.rgb = color(fill)
    if border:
        item.line.color.rgb = color(border)
        item.line.width = Pt(1.2)
    else:
        item.line.fill.background()
    return item


def arrow(slide, x, y, w=0.3, fill=TEAL):
    return polygon(slide, [(x, y), (x+w, y+0.13), (x, y+0.26)], fill)


def tag(slide, value, x, y, w, fill=TEAL, fg=WHITE):
    shape(slide, x, y, w, 0.35, fill, radius=True)
    text(slide, value, x+0.1, y+0.035, w-0.2, 0.26, 10.5, fg, True)


def iso(slide, origin, dims, colors=(TEAL, '0A6875', MINT), scale=0.006):
    """Editable isometric cuboid; dimensions and position use millimetres."""
    ox, oy = origin
    length, width, height = dims
    def p(x, y, z):
        return ox+(x-y)*scale, oy+(x+y)*scale*0.43-z*scale
    a, b, c, d = p(0,0,0), p(length,0,0), p(length,width,0), p(0,width,0)
    aa, bb, cc, dd = p(0,0,height), p(length,0,height), p(length,width,height), p(0,width,height)
    polygon(slide, [a,b,bb,aa], colors[0])
    polygon(slide, [b,c,cc,bb], colors[1])
    polygon(slide, [aa,bb,cc,dd], colors[2])
    return p


def base(number, section, title, subtitle=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(NAVY if dark else BG)
    fg = WHITE if dark else INK
    secondary = 'B1C2CA' if dark else MUTED
    text(slide, f'{number:02d} / {section.upper()}', .6, .35, 9, .25, 10.5, MINT if dark else TEAL, True)
    if title:
        text(slide, title, .6, .95, 12.05, 1.32, 33, fg, True, spacing=1.06)
    if subtitle:
        text(slide, subtitle, .6, 2.13, 12, .56, 17, secondary)
    line(slide, .6, 7.08, 12.73, 7.08, '314957' if dark else LINE, .7)
    text(slide, 'DunCarBox', .6, 7.17, 3, .2, 10, secondary, True)
    text(slide, 'ХАКАТОН / 2026', 5.3, 7.17, 3, .2, 9, secondary, align=PP_ALIGN.CENTER)
    text(slide, f'{number:02d}', 12.13, 7.15, .6, .24, 10, secondary, align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, body, sources):
    slide.notes_slide.notes_text_frame.text = body + '\n\nИсточники в репозитории:\n' + '\n'.join(sources)


def card(slide, x, y, w, h, heading, body, accent=TEAL, number=None):
    shape(slide, x, y, w, h, WHITE, radius=True)
    shape(slide, x, y, .055, h, accent)
    offset = .28
    if number:
        text(slide, number, x+.28, y+.22, w-.55, .5, 32, accent, True)
        offset = .95
    text(slide, heading, x+.28, y+offset, w-.55, .67, 21, INK, True)
    text(slide, body, x+.28, y+offset+.82, w-.55, h-offset-.96, 16, MUTED)


# 01 — opening
s = base(1, 'От заказа к действию', '', dark=True)
tag(s, 'ИНТЕЛЛЕКТУАЛЬНАЯ УПАКОВКА ЗАКАЗОВ', .6, 1.05, 4.25, '254553', MINT)
text(s, 'DunCarBox', .6, 1.74, 7.4, .9, 57, WHITE, True)
text(s, 'От заказа — к понятной\nукладке в коробки', .65, 2.98, 7.3, 1.58, 33, WHITE, True)
text(s, 'Подобрать коробки. Разместить товары.\nПоказать каждый шаг сборки.', .66, 4.98, 7.0, .88, 20, 'B1C2CA')
iso(s, (9.45, 4.75), (300, 230, 55), ('416373','304D5E','587888'), .0065)
iso(s, (9.52, 4.3), (150, 150, 220), (TEAL, '06636E', MINT), .0065)
iso(s, (10.58, 4.76), (100, 155, 135), (CORAL, 'C65946','F6AC95'), .0065)
iso(s, (9.47, 5.1), (125, 90, 80), (GOLD,'C48F3E','F5D399'), .0065)
text(s, 'ГЕОМЕТРИЯ → ПЛАН → ИНСТРУКЦИЯ', 8.0, 6.18, 4.7, .3, 11, MINT, True, PP_ALIGN.CENTER)
notes(s, '≈35 секунд.\nDunCarBox — рабочее место упаковщика. Мы берём заказ с размерами и весом товаров, доступный каталог коробок и превращаем их в конкретный план укладки. Результат отвечает сразу на три вопроса: какие коробки взять, где расположить каждый товар и в какой последовательности действовать. В проекте есть серверный расчёт, визуализация и объяснение проблемных случаев. Сегодня покажу путь от входных данных до понятной инструкции и обозначу границы модели.', ['README.md', 'docs/UX.md'])

# 02 — user problem
s = base(2, 'Задача склада', 'Коробку мало выбрать.\nНужно объяснить, как уложить.', 'При сборке одного заказа оператор решает несколько связанных задач.')
card(s, .6, 3.0, 3.86, 2.88, 'Что взять?', 'Внутренние размеры коробок,\nдопустимый вес и остаток.', TEAL, '01')
card(s, 4.73, 3.0, 3.86, 2.88, 'Как уложить?', 'Габариты, повороты, опора\nи порядок размещения.', CORAL, '02')
card(s, 8.87, 3.0, 3.86, 2.88, 'Что не вошло?', 'Сохранить готовую часть\nи объяснить остаток.', TEAL, '03')
shape(s, .6, 6.14, 12.13, .61, PALE_TEAL, radius=True)
text(s, 'Цель: сделать решение воспроизводимым и удобным для выполнения.', .87, 6.27, 11.58, .33, 18, TEAL, True)
notes(s, '≈35 секунд.\nОператору недостаточно знать суммарный объём заказа. Предметы могут не пройти по одному измерению, превысить допустимую массу или потребовать коробки, которых нет на складе. Даже корректная комбинация коробок ещё не объясняет, как именно расположить товары. Поэтому задача продукта шире подбора размера: нужно связать геометрическое решение с действиями человека. Ожидаемая ценность — более воспроизводимая сборка; экономию времени и материалов мы предлагаем измерять на реальных заказах, а не заявляем заранее.', ['docs/CONTRACTS.md', 'docs/PACKING_ENGINE.md'])

# 03 — workflow
s = base(3, 'Сценарий', 'Один путь от заказа до готовой коробки', 'Заказ и складские ограничения остаются связаны с каждым шагом результата.')
workflow = [
    ('Заказ', 'Товары, количество,\nразмеры и вес'),
    ('Расчёт', 'Эвристика или Z3;\nпрогресс и отмена'),
    ('Укладка', 'Коробки, 3D, слои\nи пошаговые действия'),
    ('Контроль', 'Неупакованные товары,\nпричины и JSON'),
]
for i, (heading, body) in enumerate(workflow):
    x = .6+i*3.08
    card(s, x, 2.92, 2.86, 2.38, heading, body, TEAL if i != 2 else CORAL)
    if i < 3:
        arrow(s, x+2.92, 3.97, .13)
shape(s, .6, 5.67, 12.13, 1.01, NAVY, radius=True)
text(s, 'Длинные списки сворачиваются', .91, 5.88, 5.24, .35, 21, WHITE, True)
text(s, 'Сводка на виду; состав заказа и остаток\nраскрываются по необходимости.', 6.35, 5.86, 6.0, .62, 17, 'C6D7DC')
notes(s, '≈40 секунд.\nСначала задаём товары и проверяем каталог коробок: внутренние размеры, предельный вес и доступное количество. Затем выбираем алгоритм и запускаем расчёт. Он идёт в фоне, интерфейс показывает этап и прошедшее время, а оператор может остановить серверную задачу. Полученный план можно просматривать по коробкам и шагам. Если заказ длинный, список товаров сворачивается; аналогично устроен список неупакованных единиц. В результате остаются понятная сводка, диагностика и полный экспорт JSON для дальнейшей работы с данными.', ['frontend/src/features', 'docs/UX.md', 'backend/app/services/packing_jobs.py'])

# 04 — model
s = base(4, 'Модель', 'Геометрия, масса, запас и опора', 'Каждый план проходит независимую проверку корректности размещений.')
shape(s, .6, 2.92, 5.24, 3.82, NAVY, radius=True)
iso(s, (2.54, 5.86), (245, 165, 115), (TEAL, '075B68', MINT), .007)
iso(s, (2.58, 4.99), (115, 130, 90), (CORAL, 'C45B49', 'F4AD98'), .007)
text(s, 'Основание товара должно\nопираться на дно или другие товары.', .9, 6.07, 4.63, .58, 16, WHITE)
rules = [
    ('Границы и пересечения', 'Товары внутри коробки и не пересекаются.'),
    ('Повороты', 'Только разрешённые для каждого товара.'),
    ('Вес и складской остаток', 'Лимит массы и доступное число коробок.'),
    ('Площадь опоры', 'Эвристика ≥ 80%; модель Z3 — 100%.'),
]
for i, (heading, body) in enumerate(rules):
    yy = 2.99+i*.83
    circle(s, 6.23, yy+.06, .18, CORAL if i == 3 else TEAL)
    text(s, heading, 6.62, yy-.02, 6.1, .34, 19, INK, True)
    text(s, body, 6.62, yy+.38, 6.1, .34, 16, MUTED)
text(s, 'Граница MVP: жёсткие прямоугольные товары; хрупкость и нагрузки пока не моделируются.', 6.23, 6.49, 6.5, .49, 11.5, MUTED)
notes(s, '≈45 секунд.\nВ модели каждая физическая единица товара — жёсткий прямоугольный параллелепипед с координатами в миллиметрах. Размещения не пересекаются и не выходят за внутренние границы коробки; учитываются разрешённые повороты, масса и остаток. Отдельное условие — опора основания: обычная эвристика допускает минимум восемьдесят процентов, Z3 требует сто процентов, включая совместную опору на несколько предметов. После расчёта независимый валидатор проверяет план. Это ещё не полная физическая симуляция: хрупкость, центр масс, прочность и траектория загрузки остаются за рамками MVP.', ['backend/app/packing/validation.py', 'docs/PACKING_ENGINE.md', 'docs/Z3_ENGINE.md'])

# 05 — algorithms
s = base(5, 'Алгоритмы', 'Два алгоритма — один формат результата')
shape(s, .6, 2.33, 5.94, 2.69, WHITE, radius=True)
shape(s, 6.79, 2.33, 5.94, 2.69, NAVY, radius=True)
tag(s, '01', .88, 2.6, .48, TEAL)
text(s, 'Быстрая эвристика', 1.55, 2.59, 4.52, .53, 25, INK, True)
text(s, 'Несколько стратегий укладки\nПроверка каждого кандидата\nБез гарантии глобального оптимума', .9, 3.42, 5.36, 1.34, 19, MUTED, spacing=1.46)
tag(s, '02', 7.08, 2.6, .48, CORAL)
text(s, 'Оптимизатор Z3', 7.75, 2.59, 4.63, .53, 25, WHITE, True)
text(s, 'Точная модель ограничений\nПоиск без лимита времени\nЗавершение поиска или ручная отмена', 7.1, 3.42, 5.32, 1.34, 19, 'C6D7DC', spacing=1.46)
text(s, 'ПРИОРИТЕТЫ ОПТИМИЗАЦИИ Z3 — ПО ПОРЯДКУ', .61, 5.37, 12, .29, 11, TEAL, True)
priorities = [('1', 'Больше\nединиц'), ('2', 'Больше объёма\nтоваров'), ('3', 'Меньше объёма\nкоробок'), ('4', 'Меньше\nкоробок')]
for i, (num, label) in enumerate(priorities):
    xx = .6+i*3.08
    circle(s, xx, 5.94, .42, PALE_TEAL)
    text(s, num, xx, 5.98, .42, .3, 15, TEAL, True, PP_ALIGN.CENTER)
    text(s, label, xx+.59, 5.89, 2.32, .71, 16.5, INK, True)
notes(s, '≈50 секунд.\nВ проекте два взаимозаменяемых движка. Эвристика строит кандидаты несколькими детерминированными стратегиями и выбирает проверенный результат. Она удобна для быстрого получения плана, но не доказывает глобальный оптимум. Z3 описывает задачу как систему ограничений с последовательными целями: сначала больше единиц, затем больше объёма товаров, затем меньше суммарный объём коробок и, наконец, меньше коробок. Ограничение по времени снято: поиск идёт до завершения или ручной отмены. Интерфейс отдельно сообщает, доказана ли оптимальность. Доказательство относится к этой математической модели, а не к неучтённым физическим свойствам товара.', ['backend/app/packing/z3_engine.py::_objective', 'backend/app/packing/z3_model.py', 'docs/Z3_ENGINE.md'])

# 06 — exact editable placement diagram
s = base(6, 'Выполнение', 'План превращается в действие', '3D, вид по слоям и инструкции используют одни и те же координаты.')
shape(s, .6, 2.9, 7.33, 3.86, WHITE, radius=True)
tag(s, 'ПРИМЕР ПО ДЕМО-ДАННЫМ', .88, 3.14, 3.1, PALE_TEAL, TEAL)
text(s, 'Коробка M · 400 × 300 × 250 мм', .89, 3.67, 6.76, .42, 19, INK, True)
# Exact top view of the first box in multiple-boxes.response.json.
bx, by, scale = 1.01, 4.13, .0074
shape(s, bx, by, 400*scale, 300*scale, BG, LINE)
demo_box = json.loads((ROOT / 'demo/multiple-boxes.response.json').read_text(encoding='utf-8'))['packed_boxes'][0]
for i, (placement, fill) in enumerate(zip(demo_box['placements'], (TEAL, CORAL, GOLD)), 1):
    pos, dim = placement['position'], placement['dimensions']
    xx = bx+pos['x']*scale
    yy = by+(300-pos['y']-dim['width'])*scale
    ww, hh = dim['length']*scale, dim['width']*scale
    shape(s, xx, yy, ww, hh, fill)
    text(s, str(i), xx, yy+.05, ww, .42, 22, WHITE, True, PP_ALIGN.CENTER)
text(s, 'Вид сверху · передняя сторона ↓', .97, 6.48, 5.8, .25, 12, MUTED)
legend = [('1', 'Чайник', TEAL), ('2', 'Фильтры', CORAL), ('3', 'Термос', GOLD)]
for i, (n,label,c) in enumerate(legend):
    yy = 4.43+i*.56
    circle(s, 5.19, yy, .25, c)
    text(s, label, 5.61, yy-.03, 2.0, .35, 17, INK)
shape(s, 8.2, 2.9, 4.53, 3.86, NAVY, radius=True)
tag(s, 'ШАГ 2 / ФИЛЬТРЫ', 8.48, 3.16, 2.15, '2F4B59', MINT)
text(s, 'Поверните и положите\nрядом с чайником', 8.49, 3.9, 3.95, 1.05, 23, WHITE, True)
text(s, 'Размер после поворота\n80 × 120 × 40 мм', 8.5, 5.07, 3.95, .68, 17, 'C6D7DC')
text(s, 'Координаты от угла коробки\nx = 220 · y = 0 · z = 0 мм', 8.5, 5.96, 3.95, .61, 15, 'C6D7DC')
notes(s, '≈45 секунд.\nСхема использует координаты первого короба из фиксированного сценария «Несколько коробок». Чайник стоит у начала координат. Набор фильтров повёрнут: после поворота его размеры восемьдесят на сто двадцать на сорок миллиметров, а начальная точка — двести двадцать по оси x. Рядом находится термос. В приложении те же данные отображаются в 3D и по слоям, а текущий предмет связан с шагом инструкции. Оператор может переключать коробки и шаги. Здесь показано основание размером четыреста на триста миллиметров; передняя сторона расположена снизу, как в интерфейсе.', ['demo/multiple-boxes.response.json', 'frontend/src/three/geometry.ts', 'backend/app/services/instructions.py'])

# 07 — partial results
partial = json.loads((ROOT / 'demo/stock-shortage.response.json').read_text(encoding='utf-8'))
s = base(7, 'Диагностика', 'Частичный результат тоже полезен', 'Сохраняем доступную укладку и объясняем, что мешает завершить заказ.')
shape(s, .6, 2.92, 5.35, 3.82, NAVY, radius=True)
tag(s, 'ДЕМО / НЕХВАТКА КОРОБОК', .89, 3.19, 3.35, '2F4B59', MINT)
text(s, f"{partial['metrics']['packed_items']} / {partial['metrics']['total_items']}", .9, 3.84, 4.52, 1.25, 72, WHITE, True)
text(s, 'наборов упаковано', .95, 5.14, 4.35, .5, 25, WHITE)
text(s, 'Одна коробка S использована.\nДля двух наборов коробок не хватает.', .96, 5.88, 4.56, .63, 16, 'C6D7DC')
text(s, 'Понятные причины', 6.4, 3.06, 6.03, .5, 24, INK, True)
causes = ['Не подходят размеры', 'Превышен допустимый вес', 'Исчерпан остаток коробок', 'Размещение не найдено при ограничениях']
for i, value in enumerate(causes):
    yy = 3.89+i*.48
    circle(s, 6.42, yy+.1, .12, CORAL if i == 2 else TEAL)
    text(s, value, 6.78, yy, 5.65, .42, 18, INK)
shape(s, 6.4, 6.17, 6.33, .57, PALE_TEAL, radius=True)
text(s, 'Список неупакованных товаров можно свернуть.', 6.61, 6.3, 5.92, .29, 14, TEAL, True)
notes(s, '≈40 секунд.\nПри нехватке коробок сервис не теряет уже найденную часть плана. В фиксированном примере из трёх подарочных наборов один помещён в единственную коробку S. Два остаются неупакованными, потому что подходящих коробок больше нет. Для других ситуаций есть отдельные причины: размеры, вес или отсутствие найденного размещения. Последнюю формулировку мы не выдаём за доказательство физической невозможности. Длинный список оставшихся товаров можно свернуть и открыть при разборе. Все единицы сохраняются в результате и экспорте JSON.', ['demo/stock-shortage.response.json', 'backend/app/packing/diagnostics.py', 'docs/UX.md'])

# 08 — architecture
s = base(8, 'Архитектура', 'Одно ядро расчёта, два способа поиска', 'Контракт размещений общий для API, алгоритмов и визуализации.')
card(s, .6, 2.93, 3.38, 1.87, 'React + TypeScript', 'Формы · Three.js · шаги\nVite', TEAL)
card(s, 4.59, 2.93, 3.78, 1.87, 'FastAPI / Python', 'Проверка запроса\nФоновые задачи и отмена', TEAL)
card(s, 8.98, 2.93, 3.75, 1.87, 'Движки упаковки', 'Эвристика / Z3\nОтдельные процессы', CORAL)
arrow(s, 4.13, 3.7, .3)
arrow(s, 8.52, 3.7, .3)
shape(s, 4.59, 5.26, 3.78, 1.31, PALE_TEAL, radius=True)
text(s, 'PostgreSQL 17', 4.88, 5.48, 3.2, .39, 21, TEAL, True)
text(s, 'Каталог коробок и остатки', 4.88, 6.0, 3.2, .32, 15, MUTED)
line(s, 6.48, 4.81, 6.48, 5.25, TEAL, 1.8)
shape(s, 8.98, 5.26, 3.75, 1.31, NAVY, radius=True)
text(s, 'Проверка + инструкции', 9.22, 5.46, 3.31, .39, 19, WHITE, True)
text(s, 'Из фактических размещений', 9.22, 5.99, 3.31, .33, 14.5, 'C6D7DC')
line(s, 10.86, 4.81, 10.86, 5.25, TEAL, 1.8)
text(s, 'РАЗВЁРТЫВАНИЕ', .63, 5.41, 3.1, .29, 11, TEAL, True)
text(s, 'Docker Compose\nСерверный контур: k3d', .63, 5.93, 3.54, .7, 18, INK)
notes(s, '≈40 секунд.\nFrontend на React и TypeScript собирает заказ и визуализирует результат с помощью Three.js. FastAPI валидирует входные данные, запускает фоновую задачу и передаёт чистые доменные модели выбранному движку. У эвристики и Z3 одинаковый формат ответа. Независимая проверка и генерация инструкций работают поверх фактических размещений. PostgreSQL хранит каталог коробок; расчёт использует снимок остатков и пока не резервирует их. Локальный стек запускается через Docker Compose, в репозитории также есть серверный контур k3d. Временные результаты фоновых задач не заменяют постоянную историю заказов.', ['docs/ARCHITECTURE.md', 'backend/app/main.py', 'frontend/package.json', 'infra/k3d-infra/README.md'])

# 09 — reproducible demo
s = base(9, 'Демонстрация', 'Четыре воспроизводимых сценария', 'Числа ниже — готовые демо-планы, а не замеры производительности алгоритмов.')
scenarios = [('simple-order', 'Простой заказ', '1 коробка'), ('multiple-boxes', 'Несколько коробок', '2 коробки'), ('oversized', 'Крупный товар', 'Не помещается'), ('stock-shortage', 'Дефицит коробок', 'Частичный план')]
for i, (sid, title, desc) in enumerate(scenarios):
    data = json.loads((ROOT / f'demo/{sid}.response.json').read_text(encoding='utf-8'))
    met = data['metrics']
    xx = .6+i*3.08
    shape(s, xx, 2.91, 2.87, 2.29, WHITE, radius=True)
    text(s, title, xx+.22, 3.17, 2.44, .57, 18, INK, True)
    text(s, f"{met['packed_items']} / {met['total_items']}", xx+.21, 3.88, 2.46, .63, 34, CORAL if i>1 else TEAL, True)
    text(s, desc, xx+.22, 4.69, 2.44, .31, 15, MUTED)
shape(s, .6, 5.59, 12.13, 1.14, NAVY, radius=True)
text(s, '10 000', .89, 5.77, 2.76, .71, 38, WHITE, True)
text(s, 'предметов', 3.29, 6.07, 1.6, .28, 15, 'C6D7DC')
text(s, '100 видов · 8 типов коробок', 5.08, 5.8, 7.15, .4, 22, WHITE, True)
text(s, 'Отдельный серверный сценарий: рассчитывается заново.', 5.11, 6.29, 7.16, .28, 14, 'C6D7DC')
notes(s, '≈40 секунд.\nДля выступления подготовлены четыре автономных сценария: простой заказ, несколько коробок, слишком крупный товар и дефицит коробок. Они воспроизводят сохранённые геометрически корректные ответы без сервера и явно помечены как демо. Числа на слайде взяты прямо из этих файлов, их нельзя интерпретировать как результат сравнения алгоритмов. Для живого расчёта выбираем серверный режим. Отдельно есть генератор заказа на десять тысяч предметов, сто видов и восемь типов коробок. Такой сценарий демонстрирует работу с большим входом, но мы не заявляем здесь доказанный оптимум Z3 или время его достижения.', ['demo/simple-order.response.json', 'demo/multiple-boxes.response.json', 'demo/oversized.response.json', 'demo/stock-shortage.response.json', 'backend/app/packing/workloads.py'])

# 10 — close / future
s = base(10, 'Дальше', 'Следующий шаг — пилот\nна реальных заказах', dark=True)
future = [
    ('01', 'Измерить', 'Время сборки, заполнение\nи ручные исправления.'),
    ('02', 'Уточнить модель', 'Хрупкость, нагрузки\nи стоимость коробок.'),
    ('03', 'Интегрировать', 'WMS, резервирование\nи история заказов.'),
]
for i, (num, heading, body) in enumerate(future):
    xx = .65+i*4.12
    text(s, num, xx, 3.04, 3.73, .65, 40, MINT if i != 1 else CORAL, True)
    text(s, heading, xx, 4.0, 3.76, .54, 25, WHITE, True)
    text(s, body, xx, 4.87, 3.76, .77, 18, 'C6D7DC')
line(s, .65, 6.06, 12.68, 6.06, '41606D', 1)
text(s, 'DunCarBox', .65, 6.35, 3.53, .51, 25, WHITE, True)
text(s, 'Понятный план для каждой коробки.', 4.52, 6.37, 8.18, .44, 24, MINT, True, PP_ALIGN.RIGHT)
notes(s, '≈40 секунд.\nСледующий этап — пилот на реальных заказах. Вначале измеряем длительность сборки, заполнение коробок и число ручных исправлений, чтобы проверить практическую пользу. Затем расширяем модель теми ограничениями, которые действительно важны складу: хрупкостью, нагрузкой и стоимостью упаковки. После этого можно связать расчёт с WMS, добавить резервирование и постоянную историю. Уже сейчас MVP соединяет заказ, геометрический план, понятные шаги и диагностику в одном интерфейсе. DunCarBox — понятный план для каждой коробки. Спасибо, готов показать сценарий или ответить на вопросы.', ['docs/TASKS.md', 'docs/PACKING_ENGINE.md', 'docs/ARCHITECTURE.md'])

prs.save(OUT)
print(f'Created {OUT.name}: {len(prs.slides)} slides, {OUT.stat().st_size:,} bytes')

# Verify native document structure and editable content after serialization.
check = Presentation(OUT)
assert len(check.slides) == 10
assert check.slide_width == Inches(W)
assert check.slide_height == Inches(H)
for index, slide in enumerate(check.slides, 1):
    assert slide.has_notes_slide, f'Slide {index}: missing speaker notes'
    assert 'Источники' in slide.notes_slide.notes_text_frame.text
    assert any(item.has_text_frame for item in slide.shapes)
print('PPTX reopened: slide size, slide count, editable text and speaker notes verified.')
